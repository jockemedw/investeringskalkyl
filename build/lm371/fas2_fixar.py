"""Fas 2: verifierade formelfixar i LM 371 (alla verifierade i formeltexten i denna mallversion).

F1  Hyresindex: kolumn E–BA på NPV rad 12–16 använde ^(år − avtalsstart). För Bef-avtal som
    startat före kalkylstart hoppade indexet från ^0 (år 1, hårdkodat) till ^(n+1) år 2.
    Nu ^(år − MAX(avtalsstart, kalkylstart)): hyran anges i kalkylstartens penningvärde och
    indexeras från året efter start-eller-kalkylstart. Oförändrat när avtalsstart ≥ kalkylstart.
F2  CA-index: IRR rad 26 hoppade från ^0 (år 1) till ^2 (år 2). Nu ^(årsindex − 1), samma
    konvention som DoU-raderna. Ändrar baslinjens IRR (dokumenterat i regression.py). Rättas i
    step4_bef.py där rad 26 genereras.
F3  Negativ exponent: Kalkyldata X23:X27 inflaterade investeringar med produktionsavslut före
    kalkylstart. Nu MAX(0, avslut − kalkylstart). Båda IF-grenarna var identiska → förenklat.
F4  LOOKUP med 2-radig sökvektor (P5:R6) i NPV rad 28 (tomträttsavgäld), D-kolumnen och
    else-grenen i alla kolumner. Nu P5:R5 som i rad 27.
F5  MROUND ger #NUM! för negativa tal (t.ex. LCC med kostnad men ingen hyra). NPV C48:E51 och
    IRR C6:C9 → ROUND(x/steg,0)*steg. Identiskt för positiva tal.
F6  #REF!-städning: IRR rad 10 D:BA (död per-årsrad), Grafer rad 54 (16 383 celler) + XFD45/53/55,
    Kalkyldata L39:N39 (viktad total kostnadshöjning återställd), flik 1 N23. Inget refererar dem.
"""
from __future__ import annotations
import re
from pathlib import Path
import openpyxl
from step4_bef import _set, KD, NPV, IRR, GRAF, FS

K = "'2. Kalkyldata'!"
SOURCE = Path(__file__).resolve().parent / "source" / "LM371_source.xlsx"


def _mround_to_round(formula: str) -> str:
    """=MROUND(expr,step) / =-MROUND(expr,step) → ROUND((expr)/step,0)*step."""
    m = re.fullmatch(r"(=-?)MROUND\((.*),(10000|100000)\)", formula)
    assert m, formula
    return f"{m.group(1)}ROUND(({m.group(2)})/{m.group(3)},0)*{m.group(3)}"


def ops_fas2() -> list[dict]:
    o: list[dict] = []
    # F1 — hyresindex år 2+ (kolumn D behåller ^0: exponenten är 0 när avtalsstart ≥ kalkylstart,
    # och för tidigare startade avtal är år 1 basåret i kalkylstartens penningvärde)
    # Range.Replace kan inte skriva MAX(…,…) (FormulaLocal) → hela formler från källan via openpyxl.
    src = openpyxl.load_workbook(SOURCE)
    npv, irr = src[NPV], src[IRR]
    for r, m in zip(range(12, 17), range(23, 28)):
        for col in range(5, 54):                                   # E..BA
            cell = npv.cell(r, col)
            old, new = f"-{K}$M{m})))", f"-MAX({K}$M{m},$D$4))))"
            assert isinstance(cell.value, str) and old in cell.value, cell.coordinate
            o.append(_set(NPV, cell.coordinate, f=cell.value.replace(old, new)))
    # F3 — negativ exponent på investeringens PV
    for r in range(23, 28):
        o.append(_set(KD, f"X{r}", f=f"=$W{r}/((1+$M$9)^MAX(0,$R{r}-$M$5))"))
    # F4 — LOOKUP-sökvektor
    for col in range(4, 54):                                       # D..BA — hela formler (Replace missade tyst)
        cell = npv.cell(28, col)
        assert isinstance(cell.value, str) and "$P$5:$R$6" in cell.value, cell.coordinate
        o.append(_set(NPV, cell.coordinate, f=cell.value.replace("$P$5:$R$6", "$P$5:$R$5")))
    # F5 — MROUND → ROUND
    for ref in ("C48", "E48", "C49", "E49", "C50", "E50", "C51", "E51"):
        # C50/E50 har redan fått M6→M11 av steg 3 — behåll det när hela formeln skrivs om från källan
        o.append(_set(NPV, ref, f=_mround_to_round(npv[ref].value).replace(f"{K}$M$6", f"{K}$M$11")))
    for ref in ("C6", "C7", "C8", "C9"):
        o.append(_set(IRR, ref, f=_mround_to_round(irr[ref].value)))
    # F6 — #REF!
    o += [{"op": "clear_contents", "sheet": IRR, "ref": "D10:BA10"},      # behåll bandets format (synlig rad)
          {"op": "clear", "sheet": GRAF, "ref": "54:54"},
          {"op": "clear", "sheet": GRAF, "ref": "XFD45"}, {"op": "clear", "sheet": GRAF, "ref": "XFD53"},
          {"op": "clear", "sheet": GRAF, "ref": "XFD55"},
          _set(KD, "L39", f="=IFERROR(SUMPRODUCT(L34:L38,$I$34:$I$38)/$I$39,0)"),
          _set(KD, "M39", f="=IFERROR(SUMPRODUCT(M34:M38,$I$34:$I$38)/$I$39,0)"),
          _set(KD, "N39", f="=IFERROR(SUMPRODUCT(N34:N38,$I$34:$I$38)/$I$39,0)"),
          {"op": "clear", "sheet": FS, "ref": "N23"}]
    return o
