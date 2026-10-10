"""Indikativ regressionsgate i LibreOffice (steg 1) — när Excel COM saknas (Linux/moln).

Kör den riktiga regression.py med comrun ersatt av en LibreOffice-motor (UNO):
    python build/lm371/lo_gate.py                 # kör gaten mot LM371_v3.xlsx
    python build/lm371/lo_gate.py <fil.xlsx>      # kör gaten mot given fil

Begränsningar (verifieras bara i Excel, steg 2): ops appliceras med openpyxl på en kopia — format,
datavalidering (probe_dv → None), utskrift och x14-dropdowns ignoreras/förstörs i kopian. LibreOffice räknar
inte om alla kolumner i Excel-datatabellen AC14:AI16; den emuleras exakt (rad 15/16 = rad 14 vid m = 0/1).
Verifierat 2026-10-09: emuleringen + omräkningen reproducerar Excels cachade v3-värden (D13/D14/D17, AD15:AI18).
"""
from __future__ import annotations
import atexit
import os
import shutil
import subprocess
import sys
import tempfile
import time
import types
import warnings
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
KD = "2. Kalkyldata"
DT_COLS = ["AD", "AE", "AF", "AG", "AH", "AI"]
CALC_OPS = {"set", "clear", "clear_contents", "merge", "unmerge"}
NOT_SUPPORTED = {"replace", "insert_rows", "datatable"}     # ändrar formler/struktur — kräver Excel


class LO:
    """En headless soffice-process via UNO-socket."""

    def __init__(self, port: int = 2002):
        import uno  # noqa: F401  (python3-uno)
        self.port = port

    def start(self):
        import uno
        self.prof = tempfile.mkdtemp()
        self.proc = subprocess.Popen(
            ["soffice", "--headless", "--norestore", "--nologo", f"-env:UserInstallation=file://{self.prof}",
             f"--accept=socket,host=localhost,port={self.port};urp;"],
            # svensk locale: tal → text i formler ("6,5 %") som i Excel på svenska
            env=dict(os.environ, LANG="sv_SE.UTF-8", LC_ALL="sv_SE.UTF-8"),
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        local = uno.getComponentContext()
        res = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        for _ in range(120):
            try:
                self.ctx = res.resolve(f"uno:socket,host=localhost,port={self.port};urp;StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(0.5)
        self.desk = self.ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", self.ctx)
        return self

    def stop(self):
        try:
            self.desk.terminate()
        except Exception:
            pass
        self.proc.wait(timeout=30)
        shutil.rmtree(self.prof, ignore_errors=True)

    def open(self, path):
        import uno
        from com.sun.star.beans import PropertyValue
        hidden = PropertyValue()
        hidden.Name, hidden.Value = "Hidden", True
        doc = self.desk.loadComponentFromURL(uno.systemPathToFileUrl(os.path.abspath(path)), "_blank", 0, (hidden,))
        doc.calculateAll()
        return doc

    @staticmethod
    def cell(doc, sheet, ref):
        return doc.Sheets.getByName(sheet).getCellRangeByName(ref)

    def get(self, doc, sheet, ref):
        c = self.cell(doc, sheet, ref)
        t = c.Type.value
        if t == "TEXT":
            return c.String
        if t == "FORMULA":
            if c.getError():
                return f"#ERR{c.getError()}"
            return c.String if c.FormulaResultType2 == 2 else c.Value     # 2 = TEXT
        return None if t == "EMPTY" else c.Value

    def set(self, doc, sheet, ref, v):
        c = self.cell(doc, sheet, ref)
        if isinstance(v, str):
            c.String = v
        else:
            c.Value = v


_engine: LO | None = None


def _lo() -> LO:
    global _engine
    if _engine is None:
        _engine = LO().start()
        atexit.register(_engine.stop)
    return _engine


def apply_openpyxl(source: Path, ops: list[dict], dest: Path) -> None:
    """Applicera de räknande ops:arna på en kopia. Format/validering/utskrift hoppas över."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(source)
    for o in ops:
        op = o["op"]
        if op in NOT_SUPPORTED:
            raise NotImplementedError(f"{op} kräver Excel COM")
        if op not in CALC_OPS:
            continue
        ws = wb[o["sheet"]]
        if op == "set":
            ws[o["ref"]] = o["formula"] if o.get("formula") is not None else o["value"]
        elif op in ("clear", "clear_contents"):
            cells = ws[o["ref"]]
            rows = cells if isinstance(cells, tuple) and isinstance(cells[0], tuple) else \
                [cells if isinstance(cells, tuple) else (cells,)]
            for row in rows:
                for c in row:
                    if not isinstance(c, openpyxl.cell.cell.MergedCell):
                        c.value = None
        elif op == "merge":
            ws.merge_cells(o["ref"])
        elif op == "unmerge" and o["ref"] in [str(r) for r in ws.merged_cells.ranges]:
            ws.unmerge_cells(o["ref"])
    wb.save(dest)


def emulate_datatable(lo: LO, doc) -> None:
    """Excel-semantik för envariabelstabellen AC14:AI16: rad 15/16 = rad 14 vid m (AD12) = 0 / 1."""
    m_orig = lo.get(doc, KD, "AD12")
    if not isinstance(m_orig, float) or lo.cell(doc, KD, "AD15").Type.value == "EMPTY":
        return                                              # ingen datatabell (källmallen)
    rows = {}
    for m, r in ((0, 15), (1, 16)):
        lo.set(doc, KD, "AD12", float(m))
        doc.calculateAll()
        rows[r] = [lo.get(doc, KD, f"{c}14") for c in DT_COLS]
    lo.set(doc, KD, "AD12", m_orig)
    for r, vals in rows.items():
        for c, v in zip(DT_COLS, vals):
            lo.set(doc, KD, f"{c}{r}", v)
    doc.calculateAll()


# ── comrun-gränssnittet ──
def run(source: Path, ops: list[dict], out: Path | None = None) -> dict:
    td = Path(tempfile.mkdtemp())
    work = td / "work.xlsx"
    try:
        apply_openpyxl(Path(source), ops, work)
        lo = _lo()
        doc = lo.open(work)
        emulate_datatable(lo, doc)
        probes, errors = {}, {}
        for o in ops:
            key = f'{o.get("sheet")}!{o.get("ref")}'
            if o["op"] == "probe":
                probes[key] = lo.get(doc, o["sheet"], o["ref"])
            elif o["op"] == "probe_dv":
                probes[key + "#dv"] = None                 # ej läsbart här — kontrolleras i Excel
            elif o["op"] == "scan":
                rng = lo.cell(doc, o["sheet"], o["ref"])
                errors[key] = sum((a.EndColumn - a.StartColumn + 1) * (a.EndRow - a.StartRow + 1)
                                  for a in rng.queryFormulaCells(4).RangeAddresses)   # FormulaResult.ERROR
        doc.close(True)
        if out is not None:
            shutil.copy2(work, out)
        return {"probes": probes, "errors": errors}
    finally:
        shutil.rmtree(td, ignore_errors=True)


def probe(sheet, ref):
    return {"op": "probe", "sheet": sheet, "ref": ref}


def probe_dv(sheet, ref):
    return {"op": "probe_dv", "sheet": sheet, "ref": ref}


def scan(sheet, ref):
    return {"op": "scan", "sheet": sheet, "ref": ref}


def main() -> int:
    sys.modules["comrun"] = types.SimpleNamespace(run=run, probe=probe, probe_dv=probe_dv, scan=scan)
    sys.path.insert(0, str(HERE))
    import regression as R
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "LM371_v3.xlsx"
    print("Baslinje källa …", end=" ", flush=True)
    R.check_source_baseline(HERE / "source" / "LM371_source.xlsx")
    print("ok")
    R.run_all(target)
    print("(LibreOffice. Med Excel: python build/lm371/build.py — kanonisk helbyggnad.)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
