"""Steg 7 (v3.1): restvärdesmodellen efter granskning — beslut A–E (2026-10-09, ANALYS_OMTAG beslut 11–15).

A  Kalkylränta restvärde frikopplad: M8 = M6 + M9 (extern yield + inflation). Objektets risk prissätts
   en gång, i kapitaliseringen (M11). Ersätter beslut 2. Känsligheten diskonteras med M8.
B  Vakansfaktorn får egen asymmetrisk skala (−0,25/0/+0,50/+1,00/+2,00 pp) och en live-kolumn
   "Motsvarar" med ekvivalent vakans: v = överskottsgrad × Δ / M11 (år N+1, fullt uthyrt).
C  Ett restvärde: flik 4 rad 84–90 definierar restvärdet vid årets slut = kapitaliseringsbart
   driftnetto år t+1 / M11. NPV (krav 1 och 3), IRR rad 7/C7 och Grafer rad 44 läser därifrån.
D  Vakansgrad (P23:P27) gäller bara år efter avtalsslut inom kalkylperioden. Kapitaliserat driftnetto
   räknas fullt uthyrt. Datavalidering spärrar fältet när avtalet täcker kalkylperioden; notis i I17.
E  Extraordinärt UH summeras bara t.o.m. horisonten (C49/E49/C78/E78). 5. IRR!B12:C12 rensas.
"""
from __future__ import annotations
from openpyxl.utils import get_column_letter, column_index_from_string
from step4_bef import (_set, _fmt, _rnd, _year_formula, YEAR_COLS, K, KD, NPV, IRR, GRAF, FS,
                       B_INT, B_VAK, B_DN_TILLK, B_REI_TILLK, B_PV_REI, N_PVDN, N_PVREI, N_RV)
from step5_kanslighet import S_OPT, S_BED, S_PES

# ── Flik 4: definitionsblock ──
R_HDR, R_VAK, R_KDN_T, R_KDN_H, R_RV_T, R_RV_H, R_BH_T = range(84, 91)
NPV_PRINT_LAST = R_BH_T + 1
IDX = {r: r - 3 for r in (R_KDN_T, R_RV_T, R_RV_H, R_BH_T)}        # HLOOKUP-index i $D$4:$BA$r
# ── Flik 2 ──
VAK_ROW = 76                                                     # "Långsiktig vakansrisk" i sektion 5
S5_ROWS = range(75, 79)
VAK_SKALA = [("Låg", -0.0025), ("Normal", 0.0), ("Förhöjd", 0.005), ("Hög", 0.01), ("Mycket hög", 0.02)]
AVTAL_OK = "AM"                                                  # hjälpkolumn: vakansgrad tillåten per objektrad
SLUTAR = "$P$5+$P$17-1"                                          # kalkylperiodens sista år
CENTER, LEFT = -4108, -4131


def _next(X: str) -> str:
    """Kolumnen till höger (D → E … BA → BB)."""
    return get_column_letter(column_index_from_string(X) + 1)


def ops_restvarde_v31() -> list[dict]:
    o: list[dict] = []

    # ── A: kalkylränta restvärde = extern yield + inflation ──
    o += [_set(KD, "M8", f="=$M$6+$M$9"),
          _set(KD, "C73", v="Extern direktavkastning (M6) är utgångsläget. Bedömningarna justerar kapitaliseringen (M11) "
                            "i alla tre kraven, ±0,25/±0,50 pp per faktor (vakansrisk −0,25…+2,00). "
                            "Kalkylränta restvärde = M6 + inflation."),
          _set(KD, "C80", v="Justerad direktavkastning (M11)"),
          ]
    for r in (S_OPT, S_BED, S_PES):
        # scenariot flyttar bara kapitaliseringen (exit-yield); diskontering = dagens krav M8
        o.append(_set(IRR, f"F{r}", f=(f"='4. NPV'!$E${N_PVDN}+'4. NPV'!$E${N_PVREI}"
                                       f"+ROUND(ROUND('4. NPV'!$E${N_RV}*{K}$M$11/C{r}/10000,0)*10000"
                                       f"/((1+{K}$M$8)^'4. NPV'!$E$47)/10000,0)*10000+'4. NPV'!$E$56")))

    # ── C + D: definitionsblock på flik 4 ──
    o += [_fmt(NPV, "B47:E47", f"B{R_HDR}:E{R_HDR}"),
          _set(NPV, f"B{R_HDR}", v="RESTVÄRDE — driftnetto år t+1 / justerad yield (M11)"),
          _fmt(NPV, "B20", f"B{R_VAK}:B{R_BH_T}"), _fmt(NPV, "C20", f"C{R_VAK}:C{R_BH_T}"),
          _fmt(NPV, "D20:BA20", f"D{R_VAK}:BA{R_BH_T}"),
          _fmt(NPV, "B29", f"B{R_RV_T}:B{R_RV_H}"), _fmt(NPV, "D29:BA29", f"D{R_RV_T}:BA{R_RV_H}"),
          _set(NPV, f"B{R_VAK}", v="Vakans tillkommande (efter avtal)"),
          _set(NPV, f"B{R_KDN_T}", v="Kapitaliseringsbart driftnetto, tillkommande"),
          _set(NPV, f"B{R_KDN_H}", v="Kapitaliseringsbart driftnetto, hela fastigheten"),
          _set(NPV, f"B{R_RV_T}", v="Restvärde tillkommande, årets slut (krav 1–2)"),
          _set(NPV, f"B{R_RV_H}", v="Restvärde hela fastigheten, årets slut (krav 3)"),
          _set(NPV, f"B{R_BH_T}", v="Bruttohyra tillkommande (ekvivalent vakans)"),
          _set(NPV, f"B{R_BH_T + 1}", v="Kapitaliseringsbart driftnetto = fullt uthyrt (vakans efter avtal återlagd), "
                                       "exkl extraordinärt UH. Restvärdet år t läser år t+1."),
          {"op": "italic", "sheet": NPV, "ref": f"B{R_BH_T + 1}"},
          {"op": "print_area", "sheet": NPV, "area": f"$B$1:$Y${NPV_PRINT_LAST}"},
          ]
    for X in YEAR_COLS:
        Y = _next(X)
        o += [_set(NPV, f"{X}{R_VAK}", f=f'=IF({X}$4="",0,{X}18-{X}{B_VAK})'),
              _set(NPV, f"{X}{R_KDN_T}", f=f'=IF({X}$4="",0,{X}{B_DN_TILLK}-{X}{B_REI_TILLK}-{X}{R_VAK})'),
              _set(NPV, f"{X}{R_KDN_H}", f=f'=IF({X}$4="",0,{X}29-{X}25-{X}18)'),
              _set(NPV, f"{X}{R_RV_T}", f=f'=IF({Y}$4="",0,{Y}{R_KDN_T}/{K}$M$11)'),
              _set(NPV, f"{X}{R_RV_H}", f=f'=IF({Y}$4="",0,{Y}{R_KDN_H}/{K}$M$11)'),
              _set(NPV, f"{X}{R_BH_T}", f=f'=IF({X}$4="",0,{X}17-{X}{B_INT})'),
              # IRR rad 7 (marknadsvärde per år) = samma definitionsrad
              _set(IRR, f"{X}7", f=f'=IF({X}$5="","",\'4. NPV\'!{X}${R_RV_T})'),
              ]
    # Grafer kolumn C ↔ NPV kolumn D (Grafer rad 44 matar diagrammen och NPV rad 34)
    for X in YEAR_COLS:
        G = get_column_letter(YEAR_COLS.index(X) + 3)
        o.append(_set(GRAF, f"{G}44", f=f"=IF({G}40=0,{G}48,'4. NPV'!{X}${R_RV_T})"))
    o.append(_set(GRAF, "B44", v="Restvärde tillkommande vid årets slut (= 4. NPV rad 88)"))

    # Krav 1 (tillkommande) och krav 3 (hela) läser restvärdet vid horisontåret
    rv_t = lambda yr: _rnd(f"HLOOKUP({yr},$D$4:$BA${R_RV_T},{IDX[R_RV_T]},FALSE)")   # noqa: E731
    rv_h = lambda yr: _rnd(f"HLOOKUP({yr},$D$4:$BA${R_RV_H},{IDX[R_RV_H]},FALSE)")   # noqa: E731
    o += [_set(NPV, f"C{N_RV}", f="=" + rv_t("C$43")), _set(NPV, f"E{N_RV}", f="=" + rv_t("D$43")),
          _set(NPV, "C50", f="=" + rv_h("C$43")), _set(NPV, "E50", f="=" + rv_h("D$43")),
          # IRR-exit = exakt samma cell som krav 1
          _set(IRR, "C7", f=f"='4. NPV'!$E${N_RV}"),
          ]

    # ── E: extraordinärt UH bara t.o.m. horisonten; vilseledande PV-cell bort ──
    o += [_set(NPV, "C49", f="=-" + _rnd("SUMIFS($D$38:$BA$38,$D$4:$BA$4,$C$44)")),
          _set(NPV, "E49", f="=-" + _rnd("SUMIFS($D$38:$BA$38,$D$4:$BA$4,$D$44)")),
          _set(NPV, f"C{N_PVREI}", f="=-" + _rnd(f"SUMIFS($D${B_PV_REI}:$BA${B_PV_REI},$D$4:$BA$4,$C$44)")),
          _set(NPV, f"E{N_PVREI}", f="=-" + _rnd(f"SUMIFS($D${B_PV_REI}:$BA${B_PV_REI},$D$4:$BA$4,$D$44)")),
          {"op": "clear_contents", "sheet": IRR, "ref": "B12:C12"},
          ]

    # ── B: vakansfaktorns skala + kolumnen "Motsvarar" ──
    o.append(_set(KD, "AK4", v="Skala långsiktig vakansrisk (pp)"))
    for i, (lbl, pp) in enumerate(VAK_SKALA):
        o += [_set(KD, f"AK{5 + i}", v=lbl), _set(KD, f"AL{5 + i}", v=pp)]
    # Motivering flyttar I:U → J:U; I = "Motsvarar"
    o += [_set(KD, "I74", v="Motsv. vakans"), _fmt(KD, "I74", "J74"), _set(KD, "J74", v="Motivering"),
          {"op": "align", "sheet": KD, "ref": "I74", "h": CENTER}]
    for r in S5_ROWS:
        o += [{"op": "unmerge", "sheet": KD, "ref": f"I{r}:U{r}"},
              _fmt(KD, f"H{r}", f"I{r}"), {"op": "clear_contents", "sheet": KD, "ref": f"I{r}"},
              {"op": "merge", "sheet": KD, "ref": f"J{r}:U{r}"}, {"op": "align", "sheet": KD, "ref": f"J{r}", "h": LEFT},
              {"op": "align", "sheet": KD, "ref": f"I{r}", "h": CENTER}]
    npv = "'4. NPV'!"
    v = (f"HLOOKUP({npv}$D$43+1,{npv}$D$4:$BA${R_KDN_T},{IDX[R_KDN_T]},FALSE)"
         f"/HLOOKUP({npv}$D$43+1,{npv}$D$4:$BA${R_BH_T},{IDX[R_BH_T]},FALSE)*H{VAK_ROW}/$M$11")
    o += [_set(KD, f"G{VAK_ROW}", v="Normal"),
          {"op": "dv_list", "sheet": KD, "ref": f"G{VAK_ROW}", "source": "=$AK$5:$AK$9",
           "prompt_title": "Långsiktig vakansrisk",
           "prompt": ("Vakans efter avtalsslut utöver vad extern yield prisar. Låg −0,25 / Normal 0 / Förhöjd +0,50 / "
                      "Hög +1,00 / Mycket hög +2,00 pp. Kolumn I visar motsvarande vakans. Extremfall: överstyr "
                      "restvärdet i 5. IRR C11.")},
          _set(KD, f"H{VAK_ROW}", f=f"=IFERROR(INDEX($AL$5:$AL$9,MATCH(G{VAK_ROW},$AK$5:$AK$9,0)),0)"),
          # "≈ 11 %" = vakans utöver vad extern yield prisar; "≈ −3 %" = lägre än marknadens antagande
          _set(KD, f"I{VAK_ROW}", f=(f'=IFERROR(IF(H{VAK_ROW}=0,"≈ 0 %",IF(H{VAK_ROW}>0,"≈ ","≈ −")'
                                     f'&ROUND(ABS({v})*100,0)&" %"),"")')),
          ]

    # ── D: vakansgradens tidsdomän ──
    o += [_fmt(KD, "O22", "P22"), _set(KD, "P22", v="efter avtal, inom kalkylperioden"),
          {"op": "wrap", "sheet": KD, "ref": "P22"},
          # mallens visning av viktad vakansgrad (flik 1 rad 88) — samma tidsdomän
          _set(KD, "K14", v="Vakans efter avtal, viktad"),
          _set(FS, "B88", v="Vakans efter avtal inom kalkylperioden (viktad)"),
          _set(KD, f"{AVTAL_OK}22", v="Vakansgrad tillåten (avtal slutar före kalkylperiodens sista år)")]
    for r in range(23, 28):
        o += [_set(KD, f"{AVTAL_OK}{r}", f=f'=OR(N{r}="",N{r}<{SLUTAR})'),
              {"op": "dv_custom", "sheet": KD, "ref": f"P{r}",
               # utan funktionsnamn/separatorer: Validation.Formula1 tolkas lokalt
               "formula": f"=${AVTAL_OK}${r}+($P${r}=0)>0",
               "err_title": "Vakans efter avtal",
               "err_msg": ("Avtalet löper hela kalkylperioden — vakansgraden får ingen effekt. "
                           "Vakansrisk efter avtalet anges i sektion 5, Långsiktig vakansrisk."),
               "prompt_title": "Vakans efter avtal",
               "prompt": ("Vakans för år efter avtalsslut inom kalkylperioden (kassaflöde). Restvärdet räknas "
                          "alltid fullt uthyrt — restvärdets vakansrisk anges i sektion 5.")}]
    o.append(_set(KD, "I17", f=(
        '=IF(OR(AND(M23<>0,M23<R23),AND(M24<>0,M24<R24),AND(M25<>0,M25<R25),AND(M26<>0,M26<R26),'
        'AND(M27<>0,M27<R27)),"Glapp, Avtalstid / Prod.avslut. ","")'
        '&IF(G13<>"","",IF(D17>$AD$20,"Höj hyran "&ROUND(D17-$AD$20,0)&" kr/kvm för krav 1–2. ",""))'
        '&IF(COUNTIF(E23:E27,"Bef")+(P9>0)>0,"Bef/bokfört värde ingår bara i krav 3. ","")'
        f'&IF(SUMPRODUCT(($P$23:$P$27>0)*({AVTAL_OK}23:{AVTAL_OK}27=FALSE))>0,'
        '"Vakansgrad utan effekt — avtalet täcker kalkylperioden (se sektion 5).","")')))
    return o


def ops_historik_v31() -> list[dict]:
    return [_fmt(FS, "R24", "R25"), _fmt(FS, "S24", "S25"),
            _set(FS, "R25", v="17. 2026-10-09 v3.1: Ett restvärde i alla krav (flik 4 rad 84–90: driftnetto år N+1, "
                             "fullt uthyrt, exkl extraordinärt UH, / justerad yield). Kalkylränta restvärde = extern "
                             "yield + inflation. Vakansfaktorn egen skala −0,25…+2,00 pp med ekvivalent vakans. "
                             "Vakansgrad bara år efter avtal inom kalkylperioden (spärr + notis). Extraordinärt UH "
                             "t.o.m. horisonten."),
            _set(FS, "S25", v="Claude/JW")]
